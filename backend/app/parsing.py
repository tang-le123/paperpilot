"""Parse PDF / DOCX / Markdown into section-aware chunks."""
from __future__ import annotations

import io
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import numpy as np
import pymupdf
from docx import Document
from PIL import Image

from app.config import (
    CHUNK_OVERLAP,
    CHUNK_SIZE,
    OCR_DPI,
    OCR_ENABLED,
    OCR_MIN_CHARS,
    PDF_PARSER,
)

_HEADING_RE = re.compile(
    r"^\s*(?:#{1,6}\s*)?"
    r"(?:\d+(?:\.\d+)*[.)]?\s*)?"
    r"(abstract|introduction|related work|background|method(?:ology)?|approach|"
    r"experiments?|results?|evaluation|discussion|conclusions?|references|bibliography|"
    r"摘要|引言|相关工作|背景|方法|实验|结果|评估|讨论|结论|参考文献)"
    r"\s*:?\s*$",
    re.IGNORECASE,
)
_CN_NUMBERED_HEADING = re.compile(r"^\s*(?:#{1,6}\s*)?[一二三四五六七八九十]+[、.．]\s*\S.{0,40}$")
_REFERENCE_MARKERS = ("references", "bibliography", "参考文献")


@dataclass
class ParsedPaper:
    title: str
    page_count: int
    chunks: list[dict]


def _is_reference_section(section: str) -> bool:
    lowered = section.lower()
    return any(marker in lowered for marker in _REFERENCE_MARKERS)


def _clean(text: str) -> str:
    text = text.replace("\u00ad", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


_OCR_ENGINE = None
_OCR_ERROR: str | None = None


def _get_ocr():
    global _OCR_ENGINE, _OCR_ERROR
    if not OCR_ENABLED:
        return None
    if _OCR_ENGINE is not None or _OCR_ERROR is not None:
        return _OCR_ENGINE
    try:
        from rapidocr_onnxruntime import RapidOCR

        _OCR_ENGINE = RapidOCR()
    except Exception as exc:
        _OCR_ERROR = str(exc)
    return _OCR_ENGINE


def _ocr_page(page: pymupdf.Page) -> str:
    engine = _get_ocr()
    if engine is None:
        return ""
    pixmap = page.get_pixmap(dpi=OCR_DPI)
    image = np.array(Image.open(io.BytesIO(pixmap.tobytes("png"))).convert("RGB"))
    result, _ = engine(image)
    if not result:
        return ""
    lines = []
    for item in result:
        if len(item) >= 2 and str(item[1]).strip():
            lines.append(str(item[1]).strip())
    return "\n".join(lines)


def _window(text: str, size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> Iterable[str]:
    text = _clean(text)
    if not text:
        return
    if len(text) <= size:
        yield text
        return
    start = 0
    while start < len(text):
        end = min(len(text), start + size)
        yield text[start:end].strip()
        if end == len(text):
            break
        start = max(end - overlap, start + 1)


def _build_chunks(
    *,
    paper_id: int,
    text: str,
    section: str,
    page: int,
    index_start: int,
) -> list[dict]:
    chunks = []
    for offset, piece in enumerate(_window(text)):
        if len(piece) < 30:
            continue
        chunks.append(
            {
                "id": f"p{paper_id}_c{index_start + offset:04d}",
                "paper_id": paper_id,
                "page": page,
                "section": section[:120],
                "content": piece,
                "chunk_index": index_start + offset,
                "is_reference": _is_reference_section(section),
            }
        )
    return chunks


def _parse_blocks(
    blocks: list[tuple[int, str]],
    *,
    paper_id: int,
    title: str,
) -> tuple[int, list[dict]]:
    """blocks: [(page, text)] -> chunks with section/page metadata."""
    chunks: list[dict] = []
    current_section = "Body"
    index = 0
    for page, page_text in blocks:
        buffer: list[str] = []

        def flush() -> None:
            nonlocal index
            body = "\n".join(buffer).strip()
            if body:
                built = _build_chunks(
                    paper_id=paper_id,
                    text=body,
                    section=current_section,
                    page=page,
                    index_start=index,
                )
                chunks.extend(built)
                index += len(built)
            buffer.clear()

        for line in page_text.splitlines():
            stripped = line.strip()
            if _HEADING_RE.match(stripped) or _CN_NUMBERED_HEADING.match(stripped):
                flush()
                current_section = stripped.lstrip("# ").strip()
            else:
                buffer.append(line)
        flush()

    return len(blocks), chunks


def parse_pdf(paper_id: int, path: Path) -> ParsedPaper:
    if PDF_PARSER == "pymupdf4llm":
        try:
            import pymupdf4llm

            page_chunks = pymupdf4llm.to_markdown(str(path), page_chunks=True)
            blocks = []
            title = path.stem
            for index, item in enumerate(page_chunks):
                metadata = item.get("metadata") or {}
                page_number = int(metadata.get("page_number") or index + 1)
                text = (item.get("text") or "").strip()
                if text:
                    blocks.append((page_number, text))
                if index == 0 and metadata.get("title"):
                    title = str(metadata["title"]).strip()
            if blocks:
                page_count, chunks = _parse_blocks(blocks, paper_id=paper_id, title=title)
                return ParsedPaper(title=title, page_count=page_count, chunks=chunks)
        except Exception:
            # Fall through to the dependency-light parser.
            pass

    document = pymupdf.open(path)
    try:
        metadata = document.metadata or {}
        title = (metadata.get("title") or "").strip() or path.stem
        blocks = []
        for page_number, page in enumerate(document, start=1):
            text = page.get_text("text", sort=True).strip()
            if OCR_ENABLED and len(text) < OCR_MIN_CHARS:
                ocr_text = _ocr_page(page)
                if len(ocr_text) > len(text):
                    text = ocr_text
            if text:
                blocks.append((page_number, text))
        if not blocks:
            raise ValueError(
                "PDF 没有可提取的文字，且 OCR 不可用或未识别到内容；"
                "请安装 rapidocr_onnxruntime，或换文本型 PDF。"
            )
        page_count, chunks = _parse_blocks(blocks, paper_id=paper_id, title=title)
        return ParsedPaper(title=title, page_count=page_count, chunks=chunks)
    finally:
        document.close()


def parse_docx(paper_id: int, path: Path) -> ParsedPaper:
    document = Document(path)
    title = (document.core_properties.title or "").strip() or path.stem
    lines: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "").lower()
        if style.startswith("heading"):
            lines.append(text)
        else:
            lines.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(" | ".join(cells))

    page_count, chunks = _parse_blocks([(1, "\n".join(lines))], paper_id=paper_id, title=title)
    return ParsedPaper(title=title, page_count=page_count, chunks=chunks)


def parse_markdown(paper_id: int, path: Path) -> ParsedPaper:
    text = path.read_text(encoding="utf-8")
    title = path.stem
    for line in text.splitlines():
        if line.startswith("# "):
            title = line[2:].strip()
            break
    page_count, chunks = _parse_blocks([(1, text)], paper_id=paper_id, title=title)
    return ParsedPaper(title=title, page_count=page_count, chunks=chunks)


def parse_file(paper_id: int, path: Path, file_type: str) -> ParsedPaper:
    suffix = file_type.lower()
    if suffix == "pdf":
        return parse_pdf(paper_id, path)
    if suffix == "docx":
        return parse_docx(paper_id, path)
    if suffix == "md":
        return parse_markdown(paper_id, path)
    raise ValueError(f"unsupported file type: {file_type}")
